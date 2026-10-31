import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_ultimate_presentation():
    prs = Presentation()
    # 16:9 widescreen format
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Professional Executive Color Palette
    NAVY_BG     = RGBColor(15, 23, 42)       # #0F172A (Deep Slate Navy)
    SLATE_CARD  = RGBColor(30, 41, 59)      # #1E293B (Dark Card)
    CYAN_ACCENT = RGBColor(14, 165, 233)     # #0EA5E9 (Vibrant Cyan)
    TEAL_ACCENT = RGBColor(20, 184, 166)     # #14B8A6 (Teal)
    GOLD_ACCENT = RGBColor(245, 158, 11)     # #F59E0B (Amber Gold)
    GREEN_SUCCESS= RGBColor(16, 185, 129)    # #10B981 (Emerald Green)
    RED_ALERT   = RGBColor(239, 68, 68)      # #EF4444 (Rose Red)
    LIGHT_BG    = RGBColor(248, 250, 252)    # #F8FAFC (Light Slate Gray)
    WHITE_CARD  = RGBColor(255, 255, 255)    # #FFFFFF (Clean Pure White)
    BORDER_GRAY = RGBColor(226, 232, 240)    # #E2E8F0 (Soft Border)
    TEXT_MAIN   = RGBColor(30, 41, 59)       # #1E293B (Dark Primary Text)
    TEXT_MUTED  = RGBColor(100, 116, 139)    # #64748B (Muted Secondary Text)
    TEXT_LIGHT  = RGBColor(241, 245, 249)    # #F1F5F9 (Light Text for Dark Cards)
    TEXT_SUBMUTED= RGBColor(148, 163, 184)   # #94A3B8 (Submuted Light)

    blank_layout = prs.slide_layouts[6]

    def set_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, subtitle_text="DATAPULSE FUSION GRC | COLLEGE PROJECT DESIGN PRESENTATION"):
        # Header banner shape
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.4), Inches(12.133), Inches(0.95))
        banner.fill.solid()
        banner.fill.fore_color.rgb = NAVY_BG
        banner.line.fill.background()

        # Left accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.4), Inches(0.18), Inches(0.95))
        bar.fill.solid()
        bar.fill.fore_color.rgb = CYAN_ACCENT
        bar.line.fill.background()

        tf = banner.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.35)
        tf.margin_top = Inches(0.1)

        p0 = tf.paragraphs[0]
        p0.text = subtitle_text.upper()
        p0.font.size = Pt(9.5)
        p0.font.bold = True
        p0.font.color.rgb = CYAN_ACCENT

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(20)
        p1.font.bold = True
        p1.font.color.rgb = WHITE_CARD

    def add_card(slide, left, top, width, height, title, items, accent_color=CYAN_ACCENT, is_dark=False, header_badge=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = SLATE_CARD if is_dark else WHITE_CARD
        card.line.color.rgb = accent_color if is_dark else BORDER_GRAY
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)

        # Title
        p0 = tf.paragraphs[0]
        if header_badge:
            r_badge = p0.add_run()
            r_badge.text = f"[{header_badge}] "
            r_badge.font.bold = True
            r_badge.font.size = Pt(11)
            r_badge.font.color.rgb = accent_color

        r_title = p0.add_run()
        r_title.text = title
        r_title.font.size = Pt(15)
        r_title.font.bold = True
        r_title.font.color.rgb = CYAN_ACCENT if is_dark else NAVY_BG

        for item in items:
            p = tf.add_paragraph()
            p.space_before = Pt(7)
            p.level = 0
            if isinstance(item, tuple):
                key, val = item
                r_key = p.add_run()
                r_key.text = "• " + key + ": "
                r_key.font.bold = True
                r_key.font.size = Pt(11.5)
                r_key.font.color.rgb = TEXT_LIGHT if is_dark else TEXT_MAIN

                r_val = p.add_run()
                r_val.text = val
                r_val.font.bold = False
                r_val.font.size = Pt(11.5)
                r_val.font.color.rgb = TEXT_SUBMUTED if is_dark else TEXT_MUTED
            else:
                r = p.add_run()
                r.text = "• " + item
                r.font.size = Pt(11.5)
                r.font.color.rgb = TEXT_LIGHT if is_dark else TEXT_MAIN

        return card

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Theme - Maximum Impact)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1, NAVY_BG)

    # Decorative Cyan Line
    dec = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.4), Inches(7.5))
    dec.fill.solid()
    dec.fill.fore_color.rgb = CYAN_ACCENT
    dec.line.fill.background()

    # Title Card Text Box
    tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "DATAPULSE FUSION GRC"
    r.font.size = Pt(40)
    r.font.bold = True
    r.font.color.rgb = CYAN_ACCENT

    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "Enterprise Governance, Risk & Compliance Management System"
    r.font.size = Pt(22)
    r.font.color.rgb = WHITE_CARD
    p.space_after = Pt(20)

    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "Project Design Presentation • Complete System Architecture, UML Diagrams & Database Design"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = TEAL_ACCENT
    p.space_after = Pt(25)

    # Key Tech Highlights Pills
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "Tech Stack: Next.js 15 | React 19 | Express 5 | Oracle APEX Cloud ORDS | Tailwind CSS | PM2"
    r.font.size = Pt(12)
    r.font.color.rgb = TEXT_SUBMUTED

    # Footer note
    tb_foot = s1.shapes.add_textbox(Inches(1.0), Inches(6.3), Inches(11.3), Inches(0.8))
    tf_foot = tb_foot.text_frame
    p = tf_foot.paragraphs[0]
    p.text = "Presented for Academic Evaluation & Project Design Review | Confidential & Proprietary"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_SUBMUTED

    s1.notes_slide.notes_text_frame.text = (
        "Good morning respected evaluators and professors. Today I am proud to present our complete project design "
        "for DataPulse Fusion GRC, an enterprise-grade Governance, Risk, and Compliance system. "
        "Over the next 15 minutes, I will walk you through the business problem, proposed solution, system architecture, "
        "module designs, UML diagrams, database schema, and live working implementation."
    )

    # =========================================================================
    # SLIDE 2: Executive Summary & Project At-A-Glance
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2, LIGHT_BG)
    add_header(s2, "Executive Summary & Project At-A-Glance")

    # 4 Quick Snapshot Metric Cards
    metrics = [
        ("3-Tier Architecture", "Next.js 15 + Express 5 + Oracle APEX ORDS Cloud Database", CYAN_ACCENT),
        ("35+ Relational Tables", "29 Master Lookup tables + 7 Transaction tables in 3NF schema", TEAL_ACCENT),
        ("Dynamic 5x5 Heatmap", "Real-time Impact (1-5) x Likelihood (1-5) risk rating computation", GOLD_ACCENT),
        ("100% API Resilience", "Bounded 8s timeouts & in-memory shadow store fallback", GREEN_SUCCESS)
    ]

    top_pos = 1.5
    for i, (m_title, m_desc, m_color) in enumerate(metrics):
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(top_pos), Inches(12.133), Inches(1.15))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE_CARD
        card.line.color.rgb = m_color
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.2)

        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = f"Key Highlight #{i+1}: {m_title} — "
        r1.font.bold = True
        r1.font.size = Pt(15)
        r1.font.color.rgb = NAVY_BG

        r2 = p.add_run()
        r2.text = m_desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = TEXT_MUTED

        top_pos += 1.35

    s2.notes_slide.notes_text_frame.text = (
        "Slide 2 gives the evaluators an immediate high-level snapshot of what makes this project exceptional: "
        "a production 3-tier stack, over 35 relational database tables, dynamic 5x5 heatmap calculations, and a resilient shadow store fallback."
    )

    # =========================================================================
    # SLIDE 3: Problem Statement vs Proposed Solution
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3, LIGHT_BG)
    add_header(s3, "Problem Statement & Proposed Solution")

    add_card(s3, 0.6, 1.5, 5.8, 5.5, "The Enterprise Problem (Pain Points)", [
        ("Siloed Risk Registers", "Organizations manage risks in disconnected Excel files with zero real-time organizational visibility."),
        ("Inconsistent Calculations", "Differing teams use arbitrary formulas for inherent vs residual risk rating without standardized matrix rules."),
        ("Lack of Audit Traceability", "Manual logs fail compliance audits due to absence of strict user attribution (X-User-*) and timestamps."),
        ("API Latency & UI Freezes", "Third-party cloud database delays (>8s) cause frontend applications to hang silently."),
        ("Keyboard Mash Data Noise", "Unsanitized user inputs introduce spam and corrupted text into risk registers.")
    ], accent_color=RED_ALERT, header_badge="CURRENT CHALLENGES")

    add_card(s3, 6.9, 1.5, 5.8, 5.5, "The DataPulse Fusion Solution", [
        ("Automated RCSA Engine", "End-to-end Risk & Control Self-Assessment workflow across Provisioning, Identification, Mitigation & Monitoring."),
        ("Real-time 5x5 Heatmap", "Dynamic interactive 5x5 grid computing inherent score (Max Impact x Max Likelihood) & residual risk."),
        ("Oracle APEX ORDS Integration", "Connected Express 5 REST API gateway with Oracle Cloud Autonomous Database."),
        ("Shadow Store Fallback", "Intercepts 8s database timeouts and returns deterministic fallback objects, guaranteeing 100% UI uptime."),
        ("Security & Anti-Junk Filter", "Middleware attaches user identity and runs regex evaluation to reject spam strings.")
    ], accent_color=GREEN_SUCCESS, header_badge="OUR PROPOSED SOLUTION")

    s3.notes_slide.notes_text_frame.text = (
        "Here we contrast the existing enterprise pain points against our proposed solution. We show judges how "
        "our platform solves real-world industry problems while maintaining bulletproof API availability."
    )

    # =========================================================================
    # SLIDE 4: System Architecture & 3-Tier Enterprise Stack
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4, LIGHT_BG)
    add_header(s4, "System Architecture & 3-Tier Enterprise Stack")

    add_card(s4, 0.6, 1.5, 3.8, 5.5, "1. Presentation Tier (Frontend)", [
        ("Core Framework", "Next.js 15 App Router & React 19"),
        ("Styling & Theme", "Tailwind CSS v4 & Lucide Icons"),
        ("Micro-Interactions", "Framer Motion animations"),
        ("State Management", "TanStack React Query v5"),
        ("Form Validation", "Zod Schemas & React Hook Form"),
        ("Live Deployment", "Hostinger Cloud (datapulsefusion.com)")
    ], accent_color=CYAN_ACCENT, header_badge="CLIENT LAYER")

    add_card(s4, 4.75, 1.5, 3.8, 5.5, "2. Application Tier (API Proxy)", [
        ("Runtime & Gateway", "Express 5.0 on Node.js runtime"),
        ("HTTP Client", "Axios with 8s bounded timeouts"),
        ("Fault Resilience", "In-memory Map shadow store"),
        ("Sanitization", "Consonant-run & vowel-ratio regex"),
        ("User Attribution", "X-User-Id, Name, Role parser"),
        ("Process Uptime", "PM2 Daemon on Hostinger VPS")
    ], accent_color=NAVY_BG, header_badge="API GATEWAY")

    add_card(s4, 8.9, 1.5, 3.8, 5.5, "3. Data Tier (Database Layer)", [
        ("Database Engine", "Oracle Cloud Autonomous DB"),
        ("REST Adapter", "Oracle APEX ORDS (REST Services)"),
        ("Auth Protocol", "OAuth2 Token / Client Credentials"),
        ("Master Tables", "29 per_mst_* lookup tables"),
        ("Transaction Tables", "7 per_trn_* transaction tables"),
        ("Schema Safety", "3NF relational schema with SQL DDL")
    ], accent_color=TEAL_ACCENT, header_badge="ORACLE CLOUD DB")

    s4.notes_slide.notes_text_frame.text = (
        "The architecture is built on a clean 3-tier enterprise pattern. Next.js 15 delivers a ultra-fast frontend, "
        "Express 5 serves as an intelligent proxy gateway with security middleware, and Oracle Autonomous DB provides enterprise data storage."
    )

    # =========================================================================
    # SLIDE 5: Detailed Design of Core Modules (RCSA Engine)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5, LIGHT_BG)
    add_header(s5, "Detailed Module Design: RCSA Lifecycle Engine")

    add_card(s5, 0.6, 1.5, 5.8, 5.5, "Module 1: RCSA Core Workflow", [
        ("Provisioning (/rcsa/provisioning)", "Initialize assessment cycles, set deadlines, and assign business unit scopes."),
        ("Identification (/rcsa/identification)", "Log risk items, root causes, causal factors, and process mapping."),
        ("Control Tagging (/rcsa/.../controls)", "Map identified risks to master controls with frequency & category tags."),
        ("Mitigation (/rcsa/mitigation)", "Build action plans, assign owners, target completion dates, and track status."),
        ("Monitoring (/rcsa/monitoring)", "Review risk levels periodically, conduct sign-offs, and lock cycles.")
    ], accent_color=CYAN_ACCENT, header_badge="RCSA PIPELINE")

    add_card(s5, 6.9, 1.5, 5.8, 5.5, "Module 2: Dynamic 5x5 Heatmap Engine", [
        ("Inherent Score Formula", "Inherent Risk = Max Impact Rating (1-5) x Max Likelihood Rating (1-5)."),
        ("Rating Matrix Bands", "Low (1-5, Green), Medium (6-11, Yellow), High (12-19, Orange), Critical (20-25, Red)."),
        ("Residual Score Formula", "Residual Risk = Inherent Score x (1 - Control Effectiveness Discount)."),
        ("Interactive Rendering", "Dynamic SVG 5x5 grid with cell counts, clickable lists, and real-time refresh."),
        ("Filtering & Drilldown", "Filter heatmaps by Business Unit, Risk Category, and Assessment Cycle.")
    ], accent_color=GOLD_ACCENT, header_badge="HEATMAP ENGINE")

    s5.notes_slide.notes_text_frame.text = (
        "Module 1 handles the end-to-end RCSA cycle: Provisioning, Risk Identification, Control Tagging, Mitigation, and Monitoring. "
        "Module 2 contains the mathematical engine that computes Inherent and Residual Risk scores on a 5x5 grid."
    )

    # =========================================================================
    # SLIDE 6: Detailed Design of Core Modules (Security & Resiliency)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6, LIGHT_BG)
    add_header(s6, "Detailed Module Design: Security & Resiliency Engine")

    add_card(s6, 0.6, 1.5, 5.8, 5.5, "Module 3: Control & KRI Engine", [
        ("Master Control Catalog", "Standardized repository of preventative, detective, and corrective controls."),
        ("Design Effectiveness", "Evaluate controls as Effective (0.4), Partially Effective (0.2), or Ineffective (0.0)."),
        ("Key Risk Indicators (KRI)", "Track quantitative metrics, set upper/lower threshold limits, and generate alerts."),
        ("Incident Logging", "Record operational loss events, financial impact, and root cause analysis.")
    ], accent_color=TEAL_ACCENT, header_badge="CONTROLS & KRI")

    add_card(s6, 6.9, 1.5, 5.8, 5.5, "Module 4: Security & Resiliency", [
        ("User Attribution Middleware", "Captures X-User-Id, Name & Role on all REST calls for full auditability (REC-07)."),
        ("Anti-Junk Input Filter", "Regex engine detects consonant runs & low vowel ratios to block spam strings (REC-10)."),
        ("Shadow Store Fallback", "In-memory Map maintains state if Oracle ORDS upstream times out (>8s) (REC-01)."),
        ("Audit Trail Logging", "Immutable record of all CREATE, UPDATE, and DELETE operations across all tables.")
    ], accent_color=NAVY_BG, header_badge="SECURITY CORE")

    s6.notes_slide.notes_text_frame.text = (
        "Module 3 governs control registers and KRI metrics. Module 4 emphasizes our security features: "
        "user attribution middleware, anti-spam sanitization, and the shadow store fallback."
    )

    # =========================================================================
    # SLIDE 7: UML Diagram 1 - Use Case Diagram
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7, LIGHT_BG)
    add_header(s7, "UML Use Case Diagram: User Roles & System Capabilities")

    add_card(s7, 0.6, 1.5, 2.8, 5.5, "Risk Administrator", [
        ("Provision Cycles", "Create & open new RCSA assessment cycles."),
        ("Master Registers", "Manage control categories, impact tables, and org units."),
        ("User & RBAC", "Assign user roles and permissions.")
    ], accent_color=NAVY_BG, header_badge="ACTOR 1")

    add_card(s7, 3.65, 1.5, 2.8, 5.5, "Risk Owner", [
        ("Identify Risks", "Log risks and link causal factors/statements."),
        ("Assess Inherent", "Rate Impact (1-5) and Likelihood (1-5)."),
        ("Tag Controls", "Link master controls & evaluate effectiveness."),
        ("Mitigation Plans", "Draft and execute action plans.")
    ], accent_color=CYAN_ACCENT, header_badge="ACTOR 2")

    add_card(s7, 6.7, 1.5, 2.8, 5.5, "Compliance Officer", [
        ("Heatmap Analytics", "View 5x5 inherent & residual heatmaps."),
        ("KRI Monitoring", "Track metrics & threshold breach alerts."),
        ("Audit Trail", "Review user action logs and system history.")
    ], accent_color=TEAL_ACCENT, header_badge="ACTOR 3")

    add_card(s7, 9.75, 1.5, 2.8, 5.5, "Executive Leader", [
        ("Dashboard KPIs", "View high-level risk posture scorecards."),
        ("Appetite Check", "Compare residual risk vs appetite limit."),
        ("Export Reports", "Download executive summary PDFs/CSVs.")
    ], accent_color=GOLD_ACCENT, header_badge="ACTOR 4")

    s7.notes_slide.notes_text_frame.text = (
        "This Use Case diagram highlights our 4 primary system personas: Risk Administrators, Risk Owners, "
        "Compliance Officers, and Executive Leaders, detailing their permissions and interactions."
    )

    # =========================================================================
    # SLIDE 8: UML Diagram 2 - Class Diagram & Data Model
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8, LIGHT_BG)
    add_header(s8, "UML Class Diagram: Core Domain Entities & Relationships")

    add_card(s8, 0.6, 1.5, 3.8, 5.5, "Core Domain Entities", [
        ("RiskItem", "uuid, riskTitle, description, categoryId, status"),
        ("InherentAssessment", "id, riskId, impactRating, likelihoodRating, inherentScore"),
        ("MasterControl", "controlId, code, title, categoryId, frequencyId, effectiveness"),
        ("MitigationPlan", "planId, riskId, actionDetails, ownerId, targetDate, status"),
        ("AuditLog", "logId, action, entity, entityId, actorId, timestamp")
    ], accent_color=NAVY_BG, header_badge="ENTITIES")

    add_card(s8, 4.75, 1.5, 3.8, 5.5, "Multiplicity & Associations", [
        ("Risk -> Assessment", "1 to 1 (Each risk has one inherent assessment)."),
        ("Risk -> Controls", "1 to N (A risk maps to multiple master controls)."),
        ("Risk -> Mitigation", "1 to N (A risk can have multiple mitigation plans)."),
        ("Cycle -> Risks", "1 to N (An RCSA cycle contains multiple risks)."),
        ("User -> AuditLog", "1 to N (User actions trigger audit entries).")
    ], accent_color=CYAN_ACCENT, header_badge="RELATIONSHIPS")

    add_card(s8, 8.9, 1.5, 3.8, 5.5, "Key Operations & Methods", [
        ("calculateInherent()", "Computes impact x likelihood rating."),
        ("calculateResidual()", "Applies control effectiveness discount."),
        ("cleanInput()", "Filters noise & keyboard-mash text."),
        ("request()", "Executes ORDS call with shadow fallback."),
        ("recordAudit()", "Writes user-attributed action to log.")
    ], accent_color=GREEN_SUCCESS, header_badge="METHODS")

    s8.notes_slide.notes_text_frame.text = (
        "The Class Diagram displays the primary domain entities, their multiplicities, and key methods such as "
        "calculateInherent(), calculateResidual(), and cleanInput()."
    )

    # =========================================================================
    # SLIDE 9: UML Diagram 3 - Activity Diagram (RCSA Workflow)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9, LIGHT_BG)
    add_header(s9, "UML Activity Diagram: RCSA Cycle Execution Flow")

    # Workflow Steps
    steps = [
        ("1. Provision Cycle", "Admin opens RCSA cycle for business unit."),
        ("2. Identify Risks", "Owner enters risk details & causal factors."),
        ("3. Rate Inherent", "Assign Impact & Likelihood (1-5 matrix)."),
        ("4. Tag Controls", "Map master controls & design effectiveness."),
        ("5. Calculate Residual", "System computes residual risk rating."),
        ("6. Mitigation Plan", "Draft action plan if residual exceeds limit."),
        ("7. Sign-off & Monitor", "Management approves cycle & monitors KRIs.")
    ]

    left_pos = 0.6
    card_width = 1.62
    gap = 0.13
    for i, (title, desc) in enumerate(steps):
        card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_pos), Inches(1.8), Inches(card_width), Inches(5.0))
        card.fill.solid()
        card.fill.fore_color.rgb = SLATE_CARD if i%2==0 else WHITE_CARD
        card.line.color.rgb = CYAN_ACCENT if i%2==0 else BORDER_GRAY
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.12)
        tf.margin_right = Inches(0.12)
        tf.margin_top = Inches(0.2)

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = CYAN_ACCENT if i%2==0 else NAVY_BG

        p1 = tf.add_paragraph()
        p1.space_before = Pt(12)
        p1.text = desc
        p1.font.size = Pt(11)
        p1.font.color.rgb = TEXT_LIGHT if i%2==0 else TEXT_MUTED

        left_pos += card_width + gap

    s9.notes_slide.notes_text_frame.text = (
        "This Activity diagram depicts the sequential workflow of an RCSA cycle, moving from cycle provisioning "
        "to risk identification, inherent assessment, control tagging, residual calculation, mitigation planning, and executive sign-off."
    )

    # =========================================================================
    # SLIDE 10: UML Diagram 4 - Sequence Diagram (API Proxy & Shadow Fallback)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_bg(s10, LIGHT_BG)
    add_header(s10, "UML Sequence Diagram: REST API Proxy & Shadow Fallback")

    add_card(s10, 0.6, 1.5, 5.8, 5.5, "Happy Path (Oracle ORDS Success <= 8s)", [
        ("1. User Form Submission", "User fills risk form on Next.js frontend."),
        ("2. Request Interceptor", "Attaches X-User-Id & X-User-Name headers."),
        ("3. Express Middleware", "Sanitizes input against keyboard-mash regex."),
        ("4. ORDS Rest Dispatch", "Forwards POST to Oracle APEX Autonomous DB."),
        ("5. DB Persistence", "Oracle executes INSERT SQL and returns 201 Created."),
        ("6. Shadow Sync", "Express updates in-memory shadow store."),
        ("7. UI Toast Refresh", "Frontend displays success toast & updates heatmap.")
    ], accent_color=GREEN_SUCCESS, header_badge="ORDS SUCCESS")

    add_card(s10, 6.9, 1.5, 5.8, 5.5, "Fallback Path (Oracle ORDS Timeout > 8s)", [
        ("1. Oracle Delay (>8s)", "Upstream database drops connection or stalls."),
        ("2. Bounded Timeout", "Express backend catches upstream timeout at 8s."),
        ("3. Shadow Store Write", "Express immediately writes payload to shadow Map store."),
        ("4. Deterministic 200 OK", "Express returns shadow object with generated UUID."),
        ("5. Zero UI Freeze", "Frontend completes request seamlessly without hanging."),
        ("6. Async Audit Log", "Audit logger records fallback event for admin review.")
    ], accent_color=GOLD_ACCENT, header_badge="SHADOW FALLBACK")

    s10.notes_slide.notes_text_frame.text = (
        "Our Sequence diagram details communication between UI, Express Proxy, Oracle ORDS, and Shadow Store. "
        "If Oracle ORDS takes longer than 8 seconds, Express gracefully switches to the shadow store fallback to prevent UI freezing."
    )

    # =========================================================================
    # SLIDE 11: Database Design - Oracle APEX ORDS Schema
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_bg(s11, LIGHT_BG)
    add_header(s11, "Database Design: Oracle APEX ORDS Relational Schema")

    add_card(s11, 0.6, 1.5, 5.8, 5.5, "Master Lookup Tables (per_mst_*)", [
        ("per_mst_business_process", "Business process definitions & sub-processes."),
        ("per_mst_risk_statement", "Standardized risk descriptions & taxonomy."),
        ("per_mst_control_register", "Master control catalog & code mapping."),
        ("per_mst_control_category", "Preventative, Detective, Corrective categories."),
        ("per_mst_control_frequency", "Daily, Weekly, Monthly, Quarterly, Annual."),
        ("per_mst_impact & likelihood", "1-5 scale rating definitions & descriptions."),
        ("per_mst_roles & user_details", "User profile and role permission mappings.")
    ], accent_color=NAVY_BG, header_badge="29 MASTER TABLES")

    add_card(s11, 6.9, 1.5, 5.8, 5.5, "Transaction Tables (per_trn_*)", [
        ("per_trn_identify_risk", "Core risk register storing identified risk instances."),
        ("per_trn_inherent_risk_assessment", "Stores impact, likelihood & inherent score."),
        ("per_trn_control_implementation", "Links risks to implemented controls."),
        ("per_trn_mitigation_plans", "Action plan details, owner, target dates."),
        ("per_trn_risk_monitoring", "Periodic review history & monitoring logs."),
        ("per_trn_impact_category", "Detailed financial/operational impact breakdown.")
    ], accent_color=CYAN_ACCENT, header_badge="7 TRANSACTION TABLES")

    s11.notes_slide.notes_text_frame.text = (
        "The relational database schema hosted on Oracle Cloud comprises over 29 Master Lookup tables (per_mst_*) "
        "and 7 Transaction tables (per_trn_*), structured for optimal 3NF relational integrity."
    )

    # =========================================================================
    # SLIDE 12: UML Diagram 5 - Risk Lifecycle State Diagram
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_bg(s12, LIGHT_BG)
    add_header(s12, "UML State Diagram: Risk Item Lifecycle States")

    states = [
        ("State 1: DRAFT", "Risk identified but missing assessment ratings. Editable by owner."),
        ("State 2: SUBMITTED", "Inherent ratings assigned (Impact x Likelihood). Awaiting control tagging."),
        ("State 3: UNDER ASSESSMENT", "Controls mapped & residual risk score computed."),
        ("State 4: MITIGATION PENDING", "Residual risk exceeds appetite limit; action plan required."),
        ("State 5: APPROVED / MONITORED", "Controls & mitigations verified. Monitored periodically.")
    ]

    top_pos = 1.5
    for title, desc in states:
        card = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(top_pos), Inches(12.133), Inches(0.95))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE_CARD
        card.line.color.rgb = CYAN_ACCENT
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.18)

        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = title + " — "
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = NAVY_BG

        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_MAIN

        top_pos += 1.1

    s12.notes_slide.notes_text_frame.text = (
        "This State diagram defines the 5 lifecycle states of a risk item: Draft, Submitted, Under Assessment, "
        "Mitigation Pending, and Approved/Monitored, showing state transition criteria."
    )

    # =========================================================================
    # SLIDE 13: User Interface & Wireframe Architecture
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_bg(s13, LIGHT_BG)
    add_header(s13, "User Interface & Wireframe Architecture")

    add_card(s13, 0.6, 1.5, 3.8, 5.5, "1. Executive Dashboard UI", [
        ("Summary KPI Cards", "Total Risks, Critical Items, Active RCSA Cycles, Health Score."),
        ("Heatmap Preview", "Compact 5x5 risk heat matrix."),
        ("Action Items", "Pending approvals and overdue mitigation tasks.")
    ], accent_color=CYAN_ACCENT, header_badge="DASHBOARD")

    add_card(s13, 4.75, 1.5, 3.8, 5.5, "2. RCSA Provisioning UI", [
        ("Cycle Manager", "Tabular view of active and archived RCSA cycles."),
        ("Provision Modal", "Form to launch cycles with date pickers & process selectors."),
        ("Status Badges", "Visual pill badges indicating Draft, In Progress, or Closed.")
    ], accent_color=TEAL_ACCENT, header_badge="PROVISIONING")

    add_card(s13, 8.9, 1.5, 3.8, 5.5, "3. Heatmap & Controls UI", [
        ("5x5 Dynamic Grid", "Color-coded cells (Green/Yellow/Orange/Red)."),
        ("Control Tagging Modal", "Search master control register & set effectiveness."),
        ("Toast Feedback", "Instant Sonner toast notifications on save/update.")
    ], accent_color=NAVY_BG, header_badge="HEATMAP & MODALS")

    s13.notes_slide.notes_text_frame.text = (
        "Our UI wireframe design focuses on sleek usability, accessibility, dark/light contrast, "
        "and clear visual hierarchy built with Tailwind CSS and Radix UI components."
    )

    # =========================================================================
    # SLIDE 14: Working Implementation & Live Demonstration Proof
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_bg(s14, LIGHT_BG)
    add_header(s14, "Working Implementation & Demonstrated Features")

    add_card(s14, 0.6, 1.5, 5.8, 5.5, "Implemented Core Capabilities", [
        ("Full CRUD Operations", "Tested CRUD across 35+ master and transaction tables."),
        ("Live 5x5 Heatmap", "Real-time calculation of Max Impact x Max Likelihood ratings."),
        ("Anti-Spam Validation", "Automatic rejection of keyboard mash strings (REC-10)."),
        ("Shadow Store Fallback", "100% API uptime during upstream DB latency (REC-01)."),
        ("User Audit Trail", "User-attributed audit history captured via custom headers (REC-07).")
    ], accent_color=GREEN_SUCCESS, header_badge="PROVEN FEATURES")

    add_card(s14, 6.9, 1.5, 5.8, 5.5, "Production Deployment & Specs", [
        ("Production Hostinger VPS", "Backend API live at api.datapulsefusion.com."),
        ("Production Frontend", "Next.js web client live at datapulsefusion.com."),
        ("PM2 Process Daemon", "Daemonized process monitoring with zero-downtime reloads."),
        ("Docker Containerization", "Dockerfile and docker-compose.yml included for containers.")
    ], accent_color=CYAN_ACCENT, header_badge="CLOUD DEPLOYMENT")

    s14.notes_slide.notes_text_frame.text = (
        "Here we summarize what has been fully built, verified, and deployed. All key QA recommendations (REC-01, REC-04, "
        "REC-07, REC-10) have been successfully implemented and tested."
    )

    # =========================================================================
    # SLIDE 15: Why Choose This Project? (Judge Appeal & Evaluation Criteria)
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    set_bg(s15, LIGHT_BG)
    add_header(s15, "Why This Project Stands Out (Judge Evaluation Rubric)")

    reasons = [
        ("1. Industry-Grade Relevance", "Solves a real-world enterprise GRC/ERM challenge faced by financial institutions and corporations."),
        ("2. Architectural Resilience", "Guarantees zero system hangs using an innovative in-memory shadow store fallback mechanism."),
        ("3. End-to-End Full-Stack Mastery", "Integrates Next.js 15, Express 5, and Oracle Autonomous DB into a seamless 3-tier system."),
        ("4. Security & Audit Excellence", "Implements strict user attribution (X-User-*) and anti-spam regex input validation."),
        ("5. Live Production Deployment", "Fully deployed on Hostinger VPS managed by PM2 process daemons and Docker containers.")
    ]

    top_pos = 1.5
    for title, desc in reasons:
        card = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(top_pos), Inches(12.133), Inches(0.95))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE_CARD
        card.line.color.rgb = GREEN_SUCCESS
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.18)

        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = title + " — "
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = NAVY_BG

        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_MAIN

        top_pos += 1.1

    s15.notes_slide.notes_text_frame.text = (
        "This slide specifically addresses the presentation judges, highlighting 5 key criteria why this project "
        "meets the highest standards of academic and engineering excellence."
    )

    # =========================================================================
    # SLIDE 16: Conclusion & Future Scope (Dark Theme)
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    set_bg(s16, NAVY_BG)

    tb = s16.shapes.add_textbox(Inches(0.6), Inches(0.6), Inches(12.133), Inches(1.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "CONCLUSION & FUTURE ROADMAP"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    add_card(s16, 0.6, 1.8, 5.8, 5.0, "Project Success Summary", [
        ("Complete GRC System", "Successfully designed and implemented an enterprise Governance, Risk & Compliance platform."),
        ("Academic Distinction", "Demonstrates software engineering best practices, 3-tier architecture, and enterprise DB design."),
        ("Bulletproof Uptime", "Guaranteed system availability through shadow store caching and bounded timeouts.")
    ], accent_color=CYAN_ACCENT, is_dark=True, header_badge="SUMMARY")

    add_card(s16, 6.9, 1.8, 5.8, 5.0, "Future Enhancements", [
        ("AI Predictive Risk Scoring", "Integrate LLM models to suggest risk ratings based on loss event descriptions."),
        ("Automated Regulatory Parsing", "NLP parser to auto-map compliance laws (GDPR, ISO 27001, SOC 2) to control registers."),
        ("Mobile Native Companion App", "React Native mobile app for field risk inspectors and mobile approvals.")
    ], accent_color=TEAL_ACCENT, is_dark=True, header_badge="ROADMAP")

    s16.notes_slide.notes_text_frame.text = (
        "In conclusion, DataPulse Fusion GRC delivers a production-ready, resilient risk management system. "
        "Future work includes AI-driven predictive risk scoring and automated compliance regulatory mapping. Thank you!"
    )

    # Save outputs
    out1 = "DataPulse_Fusion_GRC_Master_Presentation.pptx"
    out2 = os.path.expanduser("~/Downloads/DataPulse_Fusion_GRC_Master_Presentation.pptx")
    prs.save(out1)
    prs.save(out2)
    print(f"Ultimate Master Presentation saved successfully to {out1} and {out2}!")

if __name__ == "__main__":
    build_ultimate_presentation()
